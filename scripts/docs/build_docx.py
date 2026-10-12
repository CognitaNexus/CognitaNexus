#!/usr/bin/env python3
"""Render repository Markdown into readable Word using bundled python-docx.
Usage: <bundled-python> scripts/docs/build_docx.py input.md output.docx
SVG images are converted by PlantUML from their source in a temporary directory.
"""
from pathlib import Path
import re,sys,tempfile,subprocess,os,shutil
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image

def visible(s):
 s=re.sub(r'!?\[([^\]]+)\]\([^)]+\)',r'\1',s)
 s=re.sub(r'<[^>]+>','',s)
 return s.replace('**','').replace('`','').replace('\\|','|')
def inline(p,text):
 for i,x in enumerate(re.split(r'(\*\*.*?\*\*|`[^`]+`)',text)):
  if not x:continue
  r=p.add_run(visible(x));r.bold=x.startswith('**')
  if x.startswith('`'):r.font.name='DejaVu Sans Mono';r.font.size=Pt(8)
def font(style,name,size):
 style.font.name=name;style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
 style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
def build(src,dst):
 src=Path(src).resolve();dst=Path(dst).resolve();lines=src.read_text().splitlines();doc=Document()
 sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(1.7);sec.left_margin=Cm(1.7);sec.right_margin=Cm(1.7)
 for n,size in [('Normal',10),('Title',23),('Heading 1',17),('Heading 2',13),('Heading 3',11),('Heading 4',10)]:font(doc.styles[n],'Noto Sans CJK SC',size)
 norm=doc.styles['Normal'];norm.paragraph_format.space_after=Pt(3);norm.paragraph_format.line_spacing=1.1
 for n in ['Heading 1','Heading 2','Heading 3','Heading 4']:
  st=doc.styles[n];st.paragraph_format.space_before=Pt(9);st.paragraph_format.space_after=Pt(4);st.paragraph_format.keep_with_next=True
 # page footer for long specification
 p=sec.footer.paragraphs[0];p.alignment=2;r=p.add_run('CognitaNexus  Review Candidate  |  ');r.font.size=Pt(8)
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
 seen=set();temp=Path(tempfile.mkdtemp(prefix='cn-word-'));i=0;code=False;imgs=0;tables=0
 try:
  while i<len(lines):
   line=lines[i].strip();i+=1
   if not line or line=='---' or line.startswith('<!--'):continue
   if line.startswith('```'):code=not code;continue
   if code:
    p=doc.add_paragraph();r=p.add_run(line);r.font.name='DejaVu Sans Mono';r.font.size=Pt(8);p.paragraph_format.space_after=Pt(1);continue
   m=re.match(r'^!\[([^\]]*)\]\(([^)]+)\)',line)
   if m:
    img=(src.parent/m[2]).resolve()
    if img in seen:doc.add_paragraph('图示复用：'+m[1]+'，见前文同图号。');continue
    seen.add(img);caption=doc.add_paragraph(m[1]);caption.paragraph_format.keep_with_next=True
    if img.suffix=='.svg':
     puml=img.parent.parent/(img.stem+'.puml')
     if not puml.exists():raise FileNotFoundError(puml)
     target=temp/(img.stem+'.puml');shutil.copyfile(puml,target);env=os.environ.copy();env.pop('DISPLAY',None)
     subprocess.run(['plantuml','-charset','UTF-8','-tpng',str(target)],env=env,check=True,capture_output=True)
     img=target.with_suffix('.png')
    with Image.open(img) as im:w,h=im.size
    # Fit full image without clipping on portrait page. Sources remain separately available.
    width=min(17.5,22.5*w/h);p=doc.add_paragraph();p.alignment=1;p.add_run().add_picture(str(img),width=Cm(width));imgs+=1
    continue
   if line.startswith('|') and i<len(lines) and re.match(r'^\|[\s:|\-]+\|$',lines[i].strip()):
    rows=[[visible(c.strip()) for c in re.split(r'(?<!\\)\|',line)[1:-1]]];i+=1
    while i<len(lines) and lines[i].strip().startswith('|'):
     rows.append([visible(c.strip()) for c in re.split(r'(?<!\\)\|',lines[i].strip())[1:-1]]);i+=1
    n=max(map(len,rows));table=doc.add_table(rows=0,cols=n);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    weights=[max(4,min(45,sum(len(r[j]) if j<len(r) else 0 for r in rows)/len(rows))) for j in range(n)];total=sum(weights)
    widths=[17.5*w/total for w in weights]
    for col,w in zip(table.columns,widths):col.width=Cm(w)
    tblpr=table._tbl.tblPr;borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
     el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    tblpr.append(borders)
    for ri,row in enumerate(rows):
     rowobj=table.add_row();cells=rowobj.cells
     rowpr=rowobj._tr.get_or_add_trPr();rowpr.append(OxmlElement('w:cantSplit'))
     for j,c in enumerate(cells):
      c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;p=c.paragraphs[0];p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.05
      r=p.add_run(row[j] if j<len(row) else '');r.font.size=Pt(8.5);r.bold=ri==0
      if ri==0:
       sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF2');c._tc.get_or_add_tcPr().append(sh)
     if ri==0:
      trpr=table.rows[-1]._tr.get_or_add_trPr();el=OxmlElement('w:tblHeader');trpr.append(el)
    tables+=1;doc.add_paragraph();continue
   m=re.match(r'^(#{1,6}) (.+)$',line)
   if m:
    level=len(m[1]);title=visible(m[2])
    # preserve written source numbering; Word styles supply navigation only
    if level==1 and not doc.paragraphs:doc.add_paragraph(title,style='Title')
    else:
     p=doc.add_paragraph(title,style='Heading '+str(min(4,max(1,level-1))))
     if level==2 and re.match(r'\d+\. ',title) and 'SRS' in dst.name:p.paragraph_format.page_break_before=True
    continue
   if line.startswith('>'):line=line.lstrip('> ').strip()
   p=doc.add_paragraph();inline(p,line)
   if re.match(r'^[-*] ',line):p.paragraph_format.left_indent=Cm(.25)
  doc.core_properties.title='CognitaNexus 软件需求规格说明书' if 'SRS' in dst.name else 'CognitaNexus 核心模块规格说明书'
  doc.core_properties.author='CognitaNexus 项目组';doc.core_properties.subject='Review Candidate 未冻结'
  
  # Remove template paragraph border decoration while retaining table grid borders.
  for tree in [doc._element,doc.styles.element]:
   for border in list(tree.iter(qn('w:pBdr'))):border.getparent().remove(border)
  doc.save(dst);print(dst,'paragraphs',len(doc.paragraphs),'tables',tables,'unique_images',imgs)
 finally:shutil.rmtree(temp)
if __name__=='__main__':build(sys.argv[1],sys.argv[2])
